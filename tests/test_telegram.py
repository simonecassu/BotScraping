import os
import tempfile

from pokebot.db import Database
from pokebot.telegram_bot import CommandHandler, TelegramCommands, _chunks


def make(index):
    db = Database(os.path.join(tempfile.mkdtemp(), "t.db"))
    return db, CommandHandler(index, db)


def test_resolve_numbers_ranges_codes_rarity_names(index):
    _, h = make(index)
    cards, unknown = h.resolve("131 149-152 c2 sir pikachu ex xyz")
    ids = {c.id for c in cards}
    assert {"me55-131", "me55-149", "me55-150", "me55-151", "me55-152", "me55c-4"} <= ids
    assert all(c.id in ids for c in index.by_id.values() if c.rarity == "Special Illustration Rare")
    assert {"me55-53", "me55-54"} <= ids  # "pikachu ex" come nome a due parole
    assert unknown == ["xyz"]


def test_add_remove_and_missing(index):
    db, h = make(index)
    r = h.handle("/aggiungi 131 132")
    assert "2 carte" in r.text and db.wanted_ids() == {"me55-131", "me55-132"}
    r = h.handle("/ho 131")
    assert db.wanted_ids() == {"me55-132"}
    assert "Articuno" in h.handle("/mancanti").text
    h.handle("/aggiungi tutte classic")
    assert sum(1 for i in db.wanted_ids() if i.startswith("me55c-")) == 30
    h.handle("/rimuovi tutte")  # solo set principale
    assert db.wanted_ids() == {i for i in db.wanted_ids() if i.startswith("me55c-")}


def test_settings_commands(index):
    db, h = make(index)
    h.handle("/soglia 60")
    h.handle("/prezzo 80")
    h.handle("/fonti vinted ebay")
    s = db.get_settings()
    assert s["lot_min_ratio"] == 0.6 and s["max_price"] == 80 and s["sources"] == ["vinted", "ebay"]
    assert h.handle("/cerca").run_search
    assert "sconosciuto" in h.handle("/boh").text
    db.mark_seen("x:1")
    assert "conferma" in h.handle("/resetvisti").text and db.is_seen("x:1")
    h.handle("/resetvisti conferma")
    assert not db.is_seen("x:1")
    assert "/mancanti" in h.handle("/aiuto").text


def test_list_is_chunked(index):
    _, h = make(index)
    text = h.handle("/lista").text
    parts = _chunks(text)
    assert len(parts) >= 2 and all(len(p) <= 3800 for p in parts)
    assert "Exeggcute" in parts[0]


class FakeClient:
    token = "x"

    def __init__(self, updates):
        self.updates = updates
        self.sent = []
        self.menu = None

    def set_my_commands(self, commands):
        self.menu = commands
        return True

    def get_me(self):
        return "fakebot"

    def webhook_info(self):
        return {"url": "", "pending_update_count": len(self.updates)}

    def delete_webhook(self):
        return True

    def get_updates(self, offset, timeout=0):
        return [u for u in self.updates if offset is None or u["update_id"] >= offset]

    def send(self, chat_id, text, buttons=None):
        self.sent.append((str(chat_id), text))
        self.last_buttons = buttons

    def send_document(self, chat_id, filename, content, caption=""):
        self.documents = getattr(self, "documents", []) + [(filename, len(content))]

    def answer_callback(self, callback_id):
        self.answered = callback_id


def test_poll_saves_chat_id_and_ignores_strangers(index, monkeypatch):
    from pokebot import config
    monkeypatch.setattr(config, "TELEGRAM_CHAT_ID", "")
    db = Database(os.path.join(tempfile.mkdtemp(), "t.db"))
    client = FakeClient([
        {"update_id": 1, "message": {"chat": {"id": 999}, "text": "/mancanti"}},   # prima di /start: ignorato
        {"update_id": 2, "message": {"chat": {"id": 42}, "text": "/start"}},
        {"update_id": 3, "message": {"chat": {"id": 999}, "text": "/aggiungi tutte"}},  # estraneo: ignorato
        {"update_id": 4, "message": {"chat": {"id": 42}, "text": "/cerca"}},
    ])
    tc = TelegramCommands(index, db, client=client)
    assert tc.poll_once() is True
    assert db.get_kv("telegram_chat_id") == "42"
    assert db.get_kv("telegram_offset") == 5
    assert [c for c, _ in client.sent] == ["42", "42"]
    assert db.wanted_ids() == set()
    # secondo giro: nessun nuovo update
    assert tc.poll_once() is False
    # menu comandi registrato una volta sola
    assert client.menu and client.menu[0][0] == "mancanti"
    from pokebot.telegram_bot import MENU_VERSION
    assert db.get_kv("telegram_menu_version") == MENU_VERSION
    client.menu = None
    tc.poll_once()
    assert client.menu is None


def test_multiline_and_interval(index):
    db, h = make(index)
    r = h.handle("/aggiungi 131\n/intervallo 30\n/cerca")
    assert r.run_search and db.wanted_ids() == {"me55-131"}
    assert db.get_settings()["interval_minutes"] == 30
    assert "30 minuti" in r.text and "1 carte" in r.text
    assert db.get_settings()["interval_minutes"] == 30
    h.handle("/intervallo 1")  # minimo 5
    assert db.get_settings()["interval_minutes"] == 5


def test_handle_payload_from_bridge(index, monkeypatch):
    from pokebot import config
    monkeypatch.setattr(config, "TELEGRAM_CHAT_ID", "")
    db = Database(os.path.join(tempfile.mkdtemp(), "t.db"))
    client = FakeClient([])
    tc = TelegramCommands(index, db, client=client)
    assert tc.handle_payload({"chat_id": 7, "text": "/mancanti"}) is False  # prima di /start: ignorato
    assert client.sent == []
    assert tc.handle_payload({"chat_id": 7, "text": "/start"}) is False
    assert db.get_kv("telegram_chat_id") == "7"
    assert tc.handle_payload({"chat_id": 7, "text": "/aggiungi 131\n/cerca"}) is True
    assert db.wanted_ids() == {"me55-131"}
    assert tc.handle_payload({"chat_id": 8, "text": "/rimuovi 131"}) is False  # estraneo
    assert db.wanted_ids() == {"me55-131"}
    assert tc.handle_payload(None) is False and tc.handle_payload({}) is False


def test_webhook_active_skips_polling(index):
    db = Database(os.path.join(tempfile.mkdtemp(), "t.db"))
    client = FakeClient([{"update_id": 1, "message": {"chat": {"id": 1}, "text": "/start"}}])
    client.webhook_info = lambda: {"url": "https://x.workers.dev/", "pending_update_count": 0}
    tc = TelegramCommands(index, db, client=client)
    assert tc.poll_once() is False
    assert client.sent == [] and db.get_kv("telegram_offset") is None


def test_notify_many_one_message_per_card_cheapest_first(index, monkeypatch):
    from pokebot import config
    from pokebot.matcher import Matcher
    from pokebot.notifier import TelegramNotifier
    from pokebot.scrapers.base import Listing

    n = TelegramNotifier(token="t", chat_id="1")
    sent = []
    monkeypatch.setattr(n, "send", lambda text, disable_preview=False: sent.append(text) or True)
    m = Matcher(index, config.DEFAULT_SETTINGS["set_keywords"])
    wanted = {"me55-145", "me55-131", "me55-130"}
    items = []
    for i in range(12):  # 12 Zorua a prezzi decrescenti: 21, 20, ... 10
        lst = Listing("vinted", f"z{i}", f"Hisuian Zorua 145/128 n.{i} <b>", f"https://v/{i}?a=1&b=2", price=21 - i, price_text=f"{21 - i} €")
        items.append((lst, m.analyze(lst.title, "", wanted)))
    lap = Listing("wallapop", "l1", "Lapras 131/128 30th", "https://w/1", price=None)
    items.append((lap, m.analyze(lap.title, "", wanted)))
    lot = Listing("wallapop", "lot1", "Lotto 30th Lapras 131/128 e Moltres 130/128", "https://w/2", price=30, price_text="30 €")
    items.append((lot, m.analyze(lot.title, "", wanted)))

    outcomes = n.notify_many(items, max_per_card=5)
    assert len(sent) == 3  # Zorua, Lapras, Lotti
    zorua = next(t for t in sent if "Hisuian Zorua 145/128" in t and "5 più economici su 12" in t)
    assert zorua.index("10 €") < zorua.index("11 €") < zorua.index("14 €") and "15 €" not in zorua
    assert "&lt;b&gt;" in zorua and "a=1&amp;b=2" in zorua
    assert any("Lapras 131/128" in t and "prezzo n.d." in t for t in sent)
    assert any("Lotti" in t and "(2/2)" in t for t in sent)
    # inviati: i 5 Zorua più economici (indici 7..11), Lapras, lotto
    assert outcomes == [False] * 7 + [True] * 5 + [True, True]


def test_max_command(index):
    db, h = make(index)
    h.handle("/max 3")
    assert db.get_settings()["max_per_card"] == 3
    assert "max 3 annunci" in h.handle("/stato").text


def test_history_folders(index):
    db, h = make(index)
    db.add_found("vinted:1", "vinted", "Zorua 145/128 A", "https://v/1", "15 €", "", "single",
                 [{"id": "me55-145", "label": "Hisuian Zorua 145/128", "sure": True}], 1.0, True)
    db.add_found("ebay:2", "ebay", "Zorua 145/128 B", "https://e/2", "8,00 €", "Roma", "single",
                 [{"id": "me55-145", "label": "Hisuian Zorua 145/128", "sure": True}], 1.0, True)
    db.add_found("wallapop:3", "wallapop", "Lotto 30th", "https://w/3", "30 €", "", "lot",
                 [{"id": "me55-145", "label": "Hisuian Zorua 145/128", "sure": True},
                  {"id": "me55-131", "label": "Lapras 131/128", "sure": True}], 1.0, False)
    idx = h.handle("/storico")
    assert "3 trovati" in idx.text and idx.buttons
    flat = [b for row in idx.buttons for b in row]
    assert ("145 Hisuian Zorua · 2", "/storico 145") in flat
    assert any(lbl.startswith("📦 Lotti") for lbl, _ in flat)
    det = h.handle("/storico 145")
    assert "eBay.it" in det.text and "Vinted" in det.text and "Wallapop" not in det.text
    assert det.text.index("eBay.it") < det.text.index("Vinted")
    assert det.buttons == [[("⬅️ Cartelle", "/storico")]]
    lots = h.handle("/storico lotti")
    assert "Lapras 131/128" in lots.text and "non inviato" in lots.text
    assert "Nessun annuncio" in h.handle("/storico 1").text
    h.handle("/storico svuota")
    assert "Nessun annuncio trovato" in h.handle("/storico").text


def test_callback_query_is_handled(index, monkeypatch):
    from pokebot import config
    monkeypatch.setattr(config, "TELEGRAM_CHAT_ID", "")
    db = Database(os.path.join(tempfile.mkdtemp(), "t.db"))
    db.set_kv("telegram_chat_id", "42")
    client = FakeClient([
        {"update_id": 1, "callback_query": {"id": "cb1", "data": "/storico", "message": {"chat": {"id": 42}}}},
    ])
    tc = TelegramCommands(index, db, client=client)
    tc.poll_once()
    assert client.answered == "cb1" and client.sent and "annuncio" in client.sent[0][1].lower()


def test_language_command_and_status_counts(index):
    db, h = make(index)
    h.handle("/lingua tutte")
    assert db.get_settings()["language"] == "tutte"
    h.handle("/lingua ita")
    assert db.get_settings()["language"] == "ita"
    run_id = db.start_run()
    db.finish_run(run_id, 10, 1, {}, {"ebay": 6, "vinted": 4})
    st = h.handle("/stato").text
    assert "lingua: ita" in st and "eBay.it 6" in st and "Vinted 4" in st


def _hist(db, card_id, label, prices, source="vinted", key="h"):
    for i, p in enumerate(prices):
        db.add_found(f"{source}:{key}{i}", source, f"{label} {i}", f"https://h/{i}", f"{p:.2f} €", "", "single",
                     [{"id": card_id, "label": label, "sure": True}], 1.0, True)


def test_prices_progress_and_settings_commands(index):
    db, h = make(index)
    db.set_wanted_bulk(["me55-131", "me55-145"], True)
    _hist(db, "me55-131", "Lapras 131/128", [20, 10, 30], "vinted", "v")
    _hist(db, "me55-131", "Lapras 131/128", [15], "ebay", "e")
    r = h.handle("/prezzi 131")
    assert "Lapras 131/128" in r.text and "Minimo <b>10,00 €</b>" in r.text and "mediana 17,50 €" in r.text
    assert "eBay.it: min 15,00 €" in r.text and "Vinted: min 10,00 €" in r.text
    assert "Nessun prezzo" in h.handle("/prezzi 145").text
    tot = h.handle("/prezzi").text
    assert "2 carte mancanti" in tot and "<b>10,00 €</b>" in tot and "1 ancora senza prezzo" in tot
    prog = h.handle("/progresso").text
    assert "189/191" in prog and "🟩" in prog and "Illustration Rare: 2/19" in prog and "10,00 €" in prog
    assert "60%" in h.handle("/affari").text
    h.handle("/affari 50")
    assert db.get_settings()["deal_pct"] == 50
    h.handle("/affari off")
    assert db.get_settings()["deal_pct"] == 0
    h.handle("/pausa")
    assert db.get_settings()["paused"] is True
    h.handle("/riprendi")
    assert db.get_settings()["paused"] is False
    h.handle("/notte 23 8")
    assert db.get_settings()["quiet_hours"] == [23, 8]
    h.handle("/notte off")
    assert db.get_settings()["quiet_hours"] is None
    h.handle("/immagini off")
    assert db.get_settings()["images"] is False
    st = h.handle("/stato").text
    assert "affari off" in st and "immagini off" in st


def test_export_sends_document(index):
    db = Database(os.path.join(tempfile.mkdtemp(), "t.db"))
    db.set_kv("telegram_chat_id", "42")
    db.set_wanted_bulk(["me55-131"], True)
    _hist(db, "me55-131", "Lapras 131/128", [12], "vinted", "x")
    client = FakeClient([{"update_id": 1, "message": {"chat": {"id": 42}, "text": "/esporta"}}])
    tc = TelegramCommands(index, db, client=client)
    tc.poll_once()
    assert client.documents and client.documents[0][0].endswith(".xlsx") and client.documents[0][1] > 5000
    # il file è un xlsx valido con i tre fogli
    from openpyxl import load_workbook
    import io
    from pokebot.export import build_workbook
    wb = load_workbook(io.BytesIO(build_workbook(index, db)))
    assert wb.sheetnames == ["Checklist", "Annunci", "Prezzi"]
    assert wb["Checklist"].max_row == 192 and wb["Annunci"].max_row == 2


def test_cerca_targeted_command(index, monkeypatch):
    from pokebot import search as search_mod
    from pokebot.scrapers.base import BaseScraper, Listing

    class Fake(BaseScraper):
        name = "fake"
        label = "Fake"

        def __init__(self, only_italy=True):
            super().__init__(only_italy=only_italy)

        def search(self, query, limit=60):
            return [Listing("fake", "1", "Lapras 131/128 30th", "https://x/1", price=12, price_text="12 €")]

    monkeypatch.setitem(search_mod.SCRAPERS, "fake", Fake)
    db, h = make(index)
    db.save_settings({"sources": ["fake"]})
    r = h.handle("/cerca 131")
    assert "in vendita adesso: 1" in r.text and "12 €" in r.text and not r.run_search
    assert r.buttons and r.buttons[0][0] == ("💶 Prezzi", "/prezzi 131")
    assert h.handle("/cerca").run_search


def test_group_uses_listing_photo_album(index, monkeypatch):
    from pokebot import config
    from pokebot.matcher import Matcher
    from pokebot.notifier import TelegramNotifier
    from pokebot.scrapers.base import Listing

    n = TelegramNotifier(token="t", chat_id="1")
    calls = []
    monkeypatch.setattr(n, "send_album", lambda photos, caption: calls.append(("album", list(photos), caption)) or True)
    import pokebot.collage as collage_mod
    monkeypatch.setattr(collage_mod, "build_collage", lambda title, rows: None)  # senza rete: niente collage, si usa l'album
    monkeypatch.setattr(n, "send_photo", lambda url, caption: calls.append(("photo", url, caption)) or True)
    monkeypatch.setattr(n, "send", lambda text, disable_preview=False: calls.append(("text", text)) or True)
    m = Matcher(index, config.DEFAULT_SETTINGS["set_keywords"])
    items = []
    for i in range(3):
        lst = Listing("vinted", str(i), f"Lapras 131/128 30th n.{i}", f"https://v/{i}", price=10 + i, price_text=f"{10 + i} €",
                      image=f"https://img/{i}.jpg")
        items.append((lst, m.analyze(lst.title, "", {"me55-131"})))
    assert n.notify_many(items, max_per_card=5, images=True) == [True] * 3
    kind, photos, caption = calls[0]
    assert kind == "album" and photos == ["https://img/0.jpg", "https://img/1.jpg", "https://img/2.jpg"]
    assert "10 €" in caption and "Lapras 131/128" in caption
    # senza foto degli annunci: immagine ufficiale della carta
    calls.clear()
    items2 = [(Listing("ebay", "x", "Lapras 131/128 30th", "https://e/x", price=5, price_text="5 €"), items[0][1])]
    n.notify_many(items2, images=True)
    assert calls[0][0] == "photo" and "scrydex" in calls[0][1]
    # immagini disattivate: solo testo
    calls.clear()
    n.notify_many(items, images=False)
    assert calls[0][0] == "text"


def test_invite_adds_second_person_and_notifier_sends_to_all(index, monkeypatch):
    from pokebot import config
    from pokebot.db import Database
    from pokebot.notifier import TelegramNotifier
    from pokebot.telegram_bot import CommandHandler, TelegramCommands

    monkeypatch.setattr(config, "TELEGRAM_CHAT_ID", "")
    db = Database(os.path.join(tempfile.mkdtemp(), "t.db"))
    tc = TelegramCommands(index, db, client=FakeClient([]))
    assert tc.handle_payload({"chat_id": 7, "text": "/start"}) is False  # proprietario
    assert tc.handle_payload({"chat_id": 8, "text": "/invita"}) is False  # estraneo: ignorato
    assert tc.handle_payload({"chat_id": 7, "text": "/invita"}) is False
    code = db.get_kv("invite_code")["code"]
    assert tc.handle_payload({"chat_id": 8, "text": "/start SBAGLIATO"}) is False
    assert db.chat_ids() == ["7"]
    assert tc.handle_payload({"chat_id": 8, "text": f"/start {code.lower()}"}) is False
    assert db.chat_ids() == ["7", "8"] and db.get_kv("invite_code") is None  # monouso
    assert tc.handle_payload({"chat_id": 8, "text": "/aggiungi 131\n/cerca"}) is True  # ora è autorizzata
    assert tc.handle_payload({"chat_id": 9, "text": f"/start {code}"}) is False and db.chat_ids() == ["7", "8"]

    h = CommandHandler(index, db)
    assert "Solo il proprietario" in h.handle("/espelli 8", "8").text
    assert "7" in h.handle("/utenti", "8").text and "proprietario" in h.handle("/utenti", "8").text
    assert "scollegata" in h.handle("/espelli 8", "7").text and db.chat_ids() == ["7"]
    assert "non si può" in h.handle("/espelli 7", "7").text

    db.add_chat_id("8")
    n = TelegramNotifier.from_db(db)
    n.token = "t"
    assert n.chat_ids == ["7", "8"] and n.chat_id == "7"
    posted = []

    class R:
        status_code = 200
        text = ""

    import pokebot.notifier as nmod
    monkeypatch.setattr(nmod.requests, "post", lambda url, **kw: posted.append((url.rsplit("/", 1)[1], kw)) or R())
    assert n.send("ciao") is True
    assert [m for m, _ in posted] == ["sendMessage", "sendMessage"]
    assert sorted(kw["json"]["chat_id"] for _, kw in posted) == ["7", "8"]


def test_lite_tier_locks_extras(index, monkeypatch):
    from pokebot import config
    from pokebot.db import Database
    from pokebot.telegram_bot import CommandHandler

    monkeypatch.setattr(config, "LITE", True)
    h = CommandHandler(index, Database(os.path.join(tempfile.mkdtemp(), "t.db")))
    assert "versione completa" in h.handle("/esporta", "1").text
    assert "versione completa" in h.handle("/invita", "1").text
    assert "Versione lite" in h.handle("/aiuto", "1").text
    assert "mancano" in h.handle("/mancanti", "1").text.lower() or "Nessuna" in h.handle("/mancanti", "1").text
