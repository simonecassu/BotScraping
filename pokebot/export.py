"""Esportazione Excel: checklist, storico annunci, prezzi."""
from __future__ import annotations

import io
import time

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import stats as pstats
from .cards import CardIndex
from .db import Database


def _autosize(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_width, max(8, width + 2))


def build_workbook(index: CardIndex, db: Database) -> bytes:
    wanted = db.wanted_ids()
    rows = db.list_found()
    wb = Workbook()

    ws = wb.active
    ws.title = "Checklist"
    ws.append(["Set", "Codice", "Numero", "Nome", "Rarità", "Mancante", "Min visto €", "Mediana €", "Annunci visti"])
    for c in index.all_cards():
        st = pstats.card_prices(rows, c).overall
        ws.append([c.set_name, index.code_of[c.id], c.number, c.name, c.rarity, "SÌ" if c.id in wanted else "",
                   st.min, st.median, st.n])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    _autosize(ws)

    ws2 = wb.create_sheet("Annunci")
    ws2.append(["Data", "Fonte", "Tipo", "Carte", "Titolo", "Prezzo", "Luogo", "Inviato", "Affare", "Link"])
    for r in rows:
        ws2.append([time.strftime("%Y-%m-%d %H:%M", time.localtime(r["created_at"])), r["source"],
                    "lotto" if r["kind"] == "lot" else "singola",
                    ", ".join(m["label"] for m in r["matched"]), r["title"], r.get("price") or "", r.get("location") or "",
                    "sì" if r.get("notified") else "no", "sì" if r.get("deal") else "", r["url"]])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    ws2.freeze_panes = "A2"
    _autosize(ws2)

    ws3 = wb.create_sheet("Prezzi")
    ws3.append(["Codice", "Carta", "Fonte", "Annunci", "Min €", "Mediana €", "Max €", "Ultimo visto"])
    for c in index.all_cards():
        cp = pstats.card_prices(rows, c)
        if not cp.overall.n:
            continue
        for src, st in sorted(cp.by_source.items()):
            ws3.append([index.code_of[c.id], c.label, pstats.SOURCE_LABELS.get(src, src), st.n, st.min, st.median, st.max,
                        time.strftime("%Y-%m-%d", time.localtime(st.last_seen)) if st.last_seen else ""])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
    ws3.freeze_panes = "A2"
    _autosize(ws3)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
